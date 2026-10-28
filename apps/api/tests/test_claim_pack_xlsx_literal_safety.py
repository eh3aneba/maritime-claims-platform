"""Synthetic Claim Pack: Excel must not execute a source quote as a formula."""

from io import BytesIO

from openpyxl import load_workbook

from app.modules.claim_packs.renderers import render_xlsx as render_base_xlsx
from app.modules.claim_packs.recovery_renderers import render_xlsx
from tests.test_claim_pack_recovery_reporting import _export_snapshot


SOURCE_LITERAL = '=HYPERLINK("https://example.invalid", "fake document link")'


def _snapshot() -> dict:
    snapshot = _export_snapshot()
    snapshot["generation_note"] = SOURCE_LITERAL
    snapshot["claim"]["vessel_name"] = SOURCE_LITERAL
    snapshot["evidence_matrix"]["rows"] = [
        {
            "fact_id": "synthetic-fact",
            "topic": "machinery.condition",
            "status": "source_review_required",
            "fact_value": SOURCE_LITERAL,
            "supporting_evidence": [{
                "document_name": SOURCE_LITERAL,
                "document_version": 1,
                "document_is_current": True,
                "source_locator_value": "page 1",
                "source_quote": SOURCE_LITERAL,
            }],
            "conflicting_evidence": [],
        },
    ]
    snapshot["approved_ai_fact_source_revalidation"] = [{
        "field_path": SOURCE_LITERAL,
        "source_document_id": "fictional-source-001",
        "reason": "source_reprocessed",
    }]
    snapshot["summary"]["approved_ai_fact_source_revalidation_count"] = 1
    snapshot["recovery_review"]["decisions"][0]["rationale"] = SOURCE_LITERAL
    return snapshot


def _assert_static_literals(data: bytes, *, extended: bool) -> None:
    book = load_workbook(BytesIO(data), data_only=False)
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", (
                        f"Executable formula in controlled sheet {sheet.title}!{cell.coordinate}"
                    )
        assert book["Pack Summary"]["B4"].value == SOURCE_LITERAL  # Vessel
        assert book["Pack Summary"]["B4"].data_type == "s"
        evidence = book["Evidence Matrix"]
        assert evidence["B2"].value == SOURCE_LITERAL
        assert evidence["B2"].data_type == "s"
        assert evidence["D2"].value == SOURCE_LITERAL
        assert evidence["D2"].data_type == "s"
        assert evidence["H2"].value == SOURCE_LITERAL
        assert evidence["H2"].data_type == "s"
        assert book["Export Manifest"]["B5"].value == SOURCE_LITERAL
        assert book["Export Manifest"]["B5"].data_type == "s"
        if extended:
            decisions = [
                row for row in book["Recovery Review"].iter_rows()
                if row[0].value == "human decision"
            ]
            assert len(decisions) == 1
            assert decisions[0][3].value == SOURCE_LITERAL
            assert decisions[0][3].data_type == "s"
            assert book["AI Fact Source Review"]["A6"].value == SOURCE_LITERAL
            assert book["AI Fact Source Review"]["A6"].data_type == "s"
    finally:
        book.close()


def test_base_claim_pack_xlsx_preserves_text_and_never_creates_formulas() -> None:
    _assert_static_literals(render_base_xlsx(_snapshot()), extended=False)


def test_governed_claim_pack_xlsx_preserves_text_and_never_creates_formulas() -> None:
    _assert_static_literals(render_xlsx(_snapshot()), extended=True)
