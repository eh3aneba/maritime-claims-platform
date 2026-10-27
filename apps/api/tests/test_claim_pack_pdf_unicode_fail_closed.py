"""Synthetic R5 guard: never silently corrupt Unicode evidence in PDF handoffs."""

from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from sqlalchemy import select

from app.modules.claim_packs.recovery_renderers import render_pdf, render_xlsx
from app.modules.claim_packs.renderers import UnsupportedClaimPackPDFTextError
from app.modules.vessels.models import Vessel
from tests.db_harness import TestingSessionLocal, client, reset_database
from tests.test_claim_pack_recovery_reporting import _export_snapshot
from tests.test_claims_api import create_orion_claim


def setup_function() -> None:
    reset_database()


def test_main_pdf_rejects_unrepresentable_persian_evidence_without_text_loss() -> None:
    snapshot = _export_snapshot()
    source = "توربوشارژر خراب نیست؛ ثبت ۱۱:۳۰ UTC"
    snapshot["evidence_matrix"]["rows"] = [{
        "fact_id": "invented-fact-only",
        "topic": "machinery.condition",
        "fact_value": "Machine report, unverified",
        "status": "source_review_required",
        "supporting_evidence": [{
            "document_name": "invented-engineer.docx",
            "document_version": 1,
            "document_is_current": True,
            "source_locator_value": "page 1",
            "source_quote": source,
        }],
        "conflicting_evidence": [],
    }]
    with pytest.raises(UnsupportedClaimPackPDFTextError) as failure:
        render_pdf(snapshot)
    assert "use XLSX" in str(failure.value)
    assert source not in str(failure.value)  # no source evidence in the error

    # Unicode is preserved in the authorized alternate review-aid format.
    workbook = load_workbook(BytesIO(render_xlsx(snapshot)), read_only=True)
    try:
        values = list(workbook["Evidence Matrix"].values)
        assert source in [cell for row in values for cell in row]
    finally:
        workbook.close()


def test_governed_appendix_also_rejects_persian_review_narratives() -> None:
    snapshot = _export_snapshot()
    source = "ثبت هشدار قطعی نیست — بررسی مستقل ضروری است"
    snapshot["recovery_review"]["decisions"][0]["rationale"] = source
    with pytest.raises(UnsupportedClaimPackPDFTextError):
        render_pdf(snapshot)
    workbook = load_workbook(BytesIO(render_xlsx(snapshot)), read_only=True)
    try:
        assert source in [
            cell
            for row in workbook["Recovery Review"].values
            for cell in row
        ]
    finally:
        workbook.close()


def test_ascii_pdf_handoff_remains_usable_with_governed_appendix() -> None:
    snapshot = _export_snapshot()
    payload = render_pdf(snapshot)
    document = PdfReader(BytesIO(payload))
    text = "\n".join(page.extract_text() or "" for page in document.pages)
    assert "MCRI-HM-TEST" in text
    assert "RECOVERY / TIME-BAR HUMAN REVIEW" in text


def test_authenticated_pdf_rejection_has_no_export_record_and_xlsx_is_available() -> None:
    result = create_orion_claim()
    claim_id = result["claim"]["id"]
    vessel_id = result["seed"]["orion"].id
    original = "MT آزمایشی — هشدار تأیید نشده نیست"
    with TestingSessionLocal() as db:
        vessel = db.scalar(select(Vessel).where(Vessel.id == vessel_id))
        assert vessel is not None
        vessel.name = original
        db.commit()

    url = f"/api/v1/claims/{claim_id}/claim-pack-exports"
    request = {"export_format": "pdf", "acknowledge_review_aid": True}
    rejected = client.post(url, json=request)
    assert rejected.status_code == 422, rejected.text
    assert "Use XLSX" in rejected.json()["detail"]
    assert original not in rejected.text

    before = client.get(url)
    assert before.status_code == 200, before.text
    assert before.json()["total"] == 0

    allowed = client.post(url, json={**request, "export_format": "xlsx"})
    assert allowed.status_code == 201, allowed.text
    assert allowed.json()["export_format"] == "xlsx"
    response = client.get(url + "/" + allowed.json()["id"] + "/download")
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        summary = dict(list(workbook["Pack Summary"].values)[1:])
        assert summary["Vessel"] == original
    finally:
        workbook.close()
