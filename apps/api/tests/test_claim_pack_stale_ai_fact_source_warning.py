"""Fictional H&M claim-pack disclosure of previously approved stale AI facts."""
from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.claim_packs.recovery_renderers import render_pdf, render_xlsx
from app.modules.claim_packs.recovery_service import build_claim_pack_snapshot
from app.modules.claims.models import Claim
from app.modules.intelligence.models import DocumentExtraction
from app.modules.processing.models import AI_SOURCE_REPROCESSED_WARNING
from app.modules.users.models import User
from tests.db_harness import TestingSessionLocal, client, reset_database
from tests.test_human_ai_review import login, seed_review_candidates


def setup_function() -> None:
    reset_database()


def _pack(ids: dict[str, str]) -> dict:
    with TestingSessionLocal() as db:
        claim = db.get(Claim, UUID(ids["claim_id"]))
        user = db.get(User, UUID(ids["alpha_user_id"]))
        assert claim is not None and user is not None
        return build_claim_pack_snapshot(
            db,
            claim=claim,
            user=user,
            generation_note="Fictional H&M evidence only",
            generated_at=datetime(2026, 10, 27, 12, 0, tzinfo=UTC),
        )


def test_claim_pack_source_drift_warning_appears_in_both_formats() -> None:
    ids = seed_review_candidates()
    login("alpha", "alpha@example.com")
    approved = client.post(f"/api/v1/ai-review/{ids['maker_id']}", json={"action": "approve"})
    assert approved.status_code == 200, approved.text

    with TestingSessionLocal() as db:
        candidate = db.get(DocumentExtraction, UUID(ids["maker_id"]))
        candidate.source_verified = False
        candidate.validation_warnings = [
            *(candidate.validation_warnings or []), AI_SOURCE_REPROCESSED_WARNING,
        ]
        db.commit()

    snapshot = _pack(ids)
    assert snapshot["summary"]["approved_ai_fact_source_revalidation_count"] == 1
    assert snapshot["summary"]["review_state"] == "attention_required"
    stale = snapshot["approved_ai_fact_source_revalidation"]
    assert len(stale) == 1
    assert stale[0]["field_path"] == "equipment.maker"
    assert stale[0]["reason"] == "source_reprocessed"
    assert stale[0]["source_document_id"] == ids["document_id"]

    pdf = PdfReader(BytesIO(render_pdf(snapshot)))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Prior human-approved AI facts" in text
    assert "source_reprocessed" in text
    assert "no automatic reversal or reapproval" in text

    book = load_workbook(BytesIO(render_xlsx(snapshot)), read_only=True)
    try:
        rows = list(book["AI Fact Source Review"].values)
        assert rows[1][1] == 1
        assert rows[5][0] == "equipment.maker"
        assert rows[5][1] == ids["document_id"]
        assert rows[5][2] == "source_reprocessed"
        summary = dict(list(book["Pack Summary"].values)[1:])
        assert summary["Previously approved AI facts requiring source revalidation"] == 1
    finally:
        book.close()


def test_healthy_ai_fact_does_not_make_claim_pack_attention_required_by_itself() -> None:
    ids = seed_review_candidates()
    login("alpha", "alpha@example.com")
    approved = client.post(f"/api/v1/ai-review/{ids['maker_id']}", json={"action": "approve"})
    assert approved.status_code == 200, approved.text
    snapshot = _pack(ids)
    assert snapshot["summary"]["approved_ai_fact_source_revalidation_count"] == 0
    assert snapshot["approved_ai_fact_source_revalidation"] == []
    book = load_workbook(BytesIO(render_xlsx(snapshot)), read_only=True)
    try:
        assert list(book["AI Fact Source Review"].values)[1][1] == 0
    finally:
        book.close()
