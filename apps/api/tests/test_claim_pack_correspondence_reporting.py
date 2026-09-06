from __future__ import annotations

import json
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.claim_packs.correspondence_snapshot import build_correspondence_snapshot
from app.modules.claim_packs.recovery_renderers import render_pdf, render_xlsx
from app.modules.claim_packs.recovery_service import SNAPSHOT_SCHEMA_VERSION
from app.modules.claims.models import Claim
from tests.db_harness import TestingSessionLocal, client, reset_database
from tests.test_claim_pack_recovery_reporting import _export_snapshot
from tests.test_claims_api import create_orion_claim


def setup_function() -> None:
    reset_database()


def _expected(item: dict) -> dict:
    return {
        "expected_state_fingerprint": item["state_fingerprint"],
        "expected_state_version": item["state_version"],
    }


def _sent_standard(claim_id: str) -> dict:
    created = client.post(
        f"/api/v1/claims/{claim_id}/correspondence",
        json={
            "direction": "outbound",
            "kind": "status_update",
            "sensitivity": "standard",
            "recipient_label": "Lead Underwriter",
            "subject": "MT ORION governed external status",
            "body": "Dear Sirs,\n\nReviewed factual status for the governed Claim Pack.\n\nKind regards,",
        },
    )
    assert created.status_code == 201, created.text
    item = created.json()
    submitted = client.post(
        f"/api/v1/claims/{claim_id}/correspondence/{item['id']}/submit",
        json=_expected(item),
    )
    assert submitted.status_code == 200, submitted.text
    approved = client.post(
        f"/api/v1/claims/{claim_id}/correspondence/{item['id']}/approve",
        json={
            "note": "Manager reviewed exact wording for external dispatch.",
            "confirm_re_review": False,
            **_expected(submitted.json()),
        },
    )
    assert approved.status_code == 200, approved.text
    approved_item = approved.json()
    sent = client.post(
        f"/api/v1/claims/{claim_id}/correspondence/{item['id']}/mark-sent",
        json={
            "confirm_sent": True,
            "channel": "email",
            "external_reference": "PACK-OUT-001",
            "expected_review_hash": approved_item["latest_review"]["review_hash"],
            **_expected(approved_item),
        },
    )
    assert sent.status_code == 200, sent.text
    return sent.json()


def test_projection_includes_only_historical_records_and_excludes_sensitive_content() -> None:
    result = create_orion_claim()
    claim_id = result["claim"]["id"]
    sent = _sent_standard(claim_id)

    privileged_subject = "PRIVILEGED-SECRET-SUBJECT-139B"
    privileged_body = "PRIVILEGED-SECRET-BODY-139B"
    privileged = client.post(
        f"/api/v1/claims/{claim_id}/correspondence",
        json={
            "direction": "internal",
            "kind": "general",
            "sensitivity": "privileged_confidential",
            "subject": privileged_subject,
            "body": privileged_body,
        },
    )
    assert privileged.status_code == 201, privileged.text

    wp_subject = "WP-SECRET-SUBJECT-139B"
    wp_body = "WP-SECRET-BODY-139B"
    without_prejudice = client.post(
        f"/api/v1/claims/{claim_id}/correspondence",
        json={
            "direction": "inbound",
            "kind": "settlement",
            "sensitivity": "without_prejudice",
            "sender_label": "External Counsel",
            "subject": wp_subject,
            "body": wp_body,
            "channel": "email",
        },
    )
    assert without_prejudice.status_code == 201, without_prejudice.text

    unsent_subject = "UNSENT-APPROVED-NOT-HISTORY-139B"
    unsent = client.post(
        f"/api/v1/claims/{claim_id}/correspondence",
        json={
            "direction": "outbound",
            "kind": "follow_up",
            "sensitivity": "standard",
            "recipient_label": "Owners",
            "subject": unsent_subject,
            "body": "This approved-but-unsent wording must not be projected as communication history.",
        },
    )
    assert unsent.status_code == 201, unsent.text

    with TestingSessionLocal() as db:
        claim = db.get(Claim, UUID(claim_id))
        assert claim is not None
        snapshot = build_correspondence_snapshot(db, claim=claim)

    assert snapshot["authority"] == "downstream_reporting_context_only"
    assert snapshot["summary"]["included_count"] == 1
    assert snapshot["summary"]["excluded_sensitive_count"] == 2
    assert snapshot["policy"]["max_records"] == 50
    assert snapshot["items"][0]["id"] == sent["id"]
    assert snapshot["items"][0]["sent_review_hash"] == sent["sent_review_hash"]
    assert snapshot["items"][0]["latest_review"]["review_hash"] == sent["sent_review_hash"]
    assert snapshot["items"][0]["latest_review"]["request_context_fingerprint"] is None

    serialized = json.dumps(snapshot, default=str)
    assert privileged_subject not in serialized
    assert privileged_body not in serialized
    assert wp_subject not in serialized
    assert wp_body not in serialized
    assert unsent_subject not in serialized
    assert "MT ORION governed external status" in serialized


def _render_snapshot() -> dict:
    snapshot = _export_snapshot()
    snapshot["snapshot_schema_version"] = "1.4"
    snapshot["correspondence_history"] = {
        "authority": "downstream_reporting_context_only",
        "disclaimer": "Human-recorded correspondence reporting only; the platform did not send these messages.",
        "confidentiality_notice": (
            "Privileged & Confidential and Without Prejudice records are excluded by default. "
            "This export does not determine legal privilege and does not waive it."
        ),
        "policy": {
            "max_records": 50,
            "max_body_excerpt_chars": 4000,
            "included_statuses": ["filed_internal", "received_external", "sent_externally"],
            "excluded_sensitive_markings": ["privileged_confidential", "without_prejudice"],
            "excluded_sensitive_count": 2,
            "omitted_for_bound_count": 0,
        },
        "summary": {
            "included_count": 1,
            "excluded_sensitive_count": 2,
            "omitted_for_bound_count": 0,
        },
        "items": [
            {
                "id": "correspondence-1",
                "direction": "outbound",
                "kind": "status_update",
                "status": "sent_externally",
                "sensitivity": "standard",
                "channel": "email",
                "sender_label": None,
                "recipient_label": "Lead Underwriter",
                "subject": "MT ORION governed external status",
                "body_excerpt": "Reviewed factual status for the governed Claim Pack.",
                "body_truncated": False,
                "external_reference": "PACK-OUT-001",
                "occurred_at": "2026-09-06T00:00:00Z",
                "sent_at": "2026-09-06T00:00:00Z",
                "created_at": "2026-09-06T00:00:00Z",
                "request_batch_id": None,
                "requirement_ids": [],
                "state_fingerprint": "a" * 64,
                "state_version": 2,
                "content_hash": "b" * 64,
                "sent_review_hash": "c" * 64,
                "latest_review": {
                    "review_number": 1,
                    "action": "approve",
                    "review_hash": "c" * 64,
                    "content_hash": "b" * 64,
                    "correspondence_state_fingerprint": "a" * 64,
                    "request_context_fingerprint": None,
                    "reviewed_at": "2026-09-06T00:00:00Z",
                },
            }
        ],
    }
    return snapshot


def test_xlsx_export_adds_governed_correspondence_sheet_and_control_notice() -> None:
    assert SNAPSHOT_SCHEMA_VERSION == "1.4"
    workbook = load_workbook(BytesIO(render_xlsx(_render_snapshot())))
    assert "Correspondence" in workbook.sheetnames
    values = [
        str(cell.value)
        for row in workbook["Correspondence"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert any("MT ORION governed external status" in value for value in values)
    assert any("excluded by default" in value.lower() for value in values)
    assert any("does not waive" in value.lower() for value in values)


def test_pdf_export_appends_governed_correspondence_reporting_context() -> None:
    rendered = PdfReader(BytesIO(render_pdf(_render_snapshot())))
    text = "\n".join(page.extract_text() or "" for page in rendered.pages)
    normalized = " ".join(text.split())
    assert "GOVERNED CORRESPONDENCE HISTORY" in normalized
    assert "MT ORION governed external status" in normalized
    assert "excluded by default" in normalized.lower()
    assert "does not waive" in normalized.lower()
